from __future__ import annotations

import os
import shutil
from io import BytesIO
from pathlib import Path
from uuid import uuid4

os.environ["MQCHAIN_AI_DATABASE_URL"] = "sqlite:///./data/test_mqchain_ai.db"
os.environ["MQCHAIN_AI_STAGED_ARTIFACT_DIR"] = "./data/test_staged_artifacts"

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.db.database import Base, SessionLocal, engine, init_db
from app.main import app
from app.models.intake import (
    AddressCandidate,
    AddressEvidence,
    ApprovedAddress,
    ApprovedAddressObservation,
    Entity,
    IntakePreview,
    SourceDocument,
    SourceJob,
)
from app.review.approval_registry import approve_candidate_groups
from app.review.snapshot_diff import create_source_snapshot, diff_source_snapshot, mark_missing_in_latest
from app.review.source_verification import record_source_verification


@pytest.fixture(autouse=True)
def reset_db() -> None:
    test_stage = Path("./data/test_staged_artifacts")
    if test_stage.exists():
        shutil.rmtree(test_stage)
    Base.metadata.drop_all(bind=engine)
    init_db()
    yield
    Base.metadata.drop_all(bind=engine)
    if test_stage.exists():
        shutil.rmtree(test_stage)


@pytest.fixture()
def client() -> TestClient:
    with TestClient(app) as test_client:
        yield test_client


def test_multisheet_manifest_assigns_sheet_metadata_and_overrides_global(client: TestClient) -> None:
    response = client.post(
        "/api/intake/upload/preview",
        files={"file": ("multi_cex.xlsx", _multi_sheet_workbook(), "application/octet-stream")},
        data={
            "source_evidence_json": '{"entity_hint":"Global","source_url":"https://global.invalid","source_origin":"Global"}',
            "requested_source_type": "excel_upload",
        },
    )
    assert response.status_code == 200, response.text
    preview = response.json()

    profiles = preview["profile"]["metadata"]["sheet_profiles"]
    assert profiles["Bybit"]["source_url"] == "https://www.bybit.com/en/app/user/proof-of-reserve"
    assert profiles["OKX"]["source_url"] == "https://www.okx.com/proof-of-reserves"
    assert "source_manifest" in preview["profile"]["skipped_sheet_names"]

    by_sheet = {candidate["source_sheet"]: candidate for candidate in preview["candidates_preview"]}
    assert by_sheet["Bybit"]["entity_name"] == "Bybit"
    assert by_sheet["Bybit"]["source_url"] == "https://www.bybit.com/en/app/user/proof-of-reserve"
    assert by_sheet["Bybit"]["raw_reference"]["sheet_entity_hint"] == "Bybit"
    assert by_sheet["Bybit"]["raw_reference"]["sheet_snapshot_date"] == "2026-04-22"
    assert by_sheet["OKX"]["entity_name"] == "OKX"
    assert by_sheet["OKX"]["raw_reference"]["source_evidence"]["source_url"] == "https://www.okx.com/proof-of-reserves"


def test_sheet_verification_applies_only_to_matching_sheet(client: TestClient) -> None:
    job_id = _run_multisheet_job(client)
    with SessionLocal() as db:
        bybit = db.scalar(select(AddressCandidate).where(AddressCandidate.source_job_id == job_id, AddressCandidate.source_sheet == "Bybit"))
        okx = db.scalar(select(AddressCandidate).where(AddressCandidate.source_job_id == job_id, AddressCandidate.source_sheet == "OKX"))
        assert bybit is not None
        assert okx is not None
        record_source_verification(
            db,
            verification_scope="source_sheet",
            verification_status="verified",
            source_trust="official_verified",
            verified_by="pytest",
            source_job_id=job_id,
            source_document_id=bybit.source_document_id,
            source_sheet="Bybit",
            entity_name="Bybit",
            source_url=bybit.source_url,
            evidence_shape=bybit.evidence_type,
        )
        db.commit()

        result = approve_candidate_groups(db, source_job_id=job_id, dry_run=False, actor="pytest")

        assert result["groups_approved"] == 1
        assert result["groups_skipped"] == 1
        assert db.scalar(select(func.count(ApprovedAddress.id))) == 1
        assert db.scalar(select(Entity.entity_name).join(ApprovedAddress, ApprovedAddress.entity_id == Entity.id)) == "Bybit"


def test_multisheet_workbook_can_approve_bybit_and_okx_separately(client: TestClient) -> None:
    job_id = _run_multisheet_job(client)
    with SessionLocal() as db:
        for candidate in db.scalars(select(AddressCandidate).where(AddressCandidate.source_job_id == job_id)).all():
            record_source_verification(
                db,
                verification_scope="source_sheet",
                verification_status="verified",
                source_trust="official_verified",
                verified_by="pytest",
                source_job_id=job_id,
                source_document_id=candidate.source_document_id,
                source_sheet=candidate.source_sheet,
                entity_name=candidate.entity_name,
                source_url=candidate.source_url,
                evidence_shape=candidate.evidence_type,
            )
        db.commit()

        result = approve_candidate_groups(db, source_job_id=job_id, dry_run=False, actor="pytest")

        assert result["groups_approved"] == 2
        assert {row[0] for row in db.execute(select(Entity.entity_name))} == {"Bybit", "OKX"}
        assert db.scalar(select(func.count(ApprovedAddress.id))) == 2


def test_monthly_snapshot_existing_address_observes_without_duplicate_approved_address() -> None:
    with SessionLocal() as db:
        job1, candidate1 = _job_with_candidate(db, entity_name="Bybit", address="0x1111111111111111111111111111111111111111")
        _verify_candidate(db, candidate1)
        snapshot1 = create_source_snapshot(db, source_job_id=job1.id, snapshot_type="por_monthly", snapshot_period="2026-04", snapshot_date="2026-04-22")
        approve_candidate_groups(db, source_job_id=job1.id, dry_run=False, actor="pytest", source_snapshot_id=snapshot1.id)

        job2, candidate2 = _job_with_candidate(db, entity_name="Bybit", address="0x1111111111111111111111111111111111111111")
        _verify_candidate(db, candidate2)
        snapshot2 = create_source_snapshot(
            db,
            source_job_id=job2.id,
            snapshot_type="por_monthly",
            snapshot_period="2026-05",
            snapshot_date="2026-05-22",
            previous_snapshot_id=snapshot1.id,
        )
        result = approve_candidate_groups(db, source_job_id=job2.id, dry_run=False, actor="pytest", source_snapshot_id=snapshot2.id)

        assert result["addresses_created"] == 0
        assert result["observations_written"] == 1
        assert db.scalar(select(func.count(ApprovedAddress.id))) == 1
        assert db.scalar(select(func.count(ApprovedAddressObservation.id))) == 2
        approved = db.scalar(select(ApprovedAddress))
        assert approved.latest_snapshot_id == snapshot2.id
        assert approved.lifecycle_status == "active"


def test_snapshot_diff_marks_missing_in_latest_without_deleting_approved_address() -> None:
    with SessionLocal() as db:
        job1, candidate1 = _job_with_candidate(db, entity_name="Bybit", address="0x1111111111111111111111111111111111111111")
        _verify_candidate(db, candidate1)
        snapshot1 = create_source_snapshot(db, source_job_id=job1.id, snapshot_type="por_monthly", snapshot_period="2026-04", snapshot_date="2026-04-22")
        approve_candidate_groups(db, source_job_id=job1.id, dry_run=False, actor="pytest", source_snapshot_id=snapshot1.id)

        job2, candidate2 = _job_with_candidate(db, entity_name="Bybit", address="0x2222222222222222222222222222222222222222")
        _verify_candidate(db, candidate2)
        snapshot2 = create_source_snapshot(
            db,
            source_job_id=job2.id,
            snapshot_type="por_monthly",
            snapshot_period="2026-05",
            snapshot_date="2026-05-22",
            previous_snapshot_id=snapshot1.id,
        )

        diff = diff_source_snapshot(db, source_job_id=job2.id, snapshot_id=snapshot2.id)
        assert diff["new_addresses"] == 1
        assert diff["missing_in_latest"] == 1

        result = mark_missing_in_latest(db, source_job_id=job2.id, source_snapshot_id=snapshot2.id, dry_run=False)

        assert result["missing_marked"] == 1
        assert db.scalar(select(func.count(ApprovedAddress.id))) == 1
        approved = db.scalar(select(ApprovedAddress))
        assert approved.lifecycle_status == "missing_in_latest"
        assert approved.metadata_json["missing_since_snapshot_id"] == snapshot2.id


def test_unapproved_candidates_are_not_export_source_for_kv_or_registry() -> None:
    with SessionLocal() as db:
        _job_with_candidate(db, entity_name="Bybit", address="0x1111111111111111111111111111111111111111")
        assert db.scalar(select(func.count(AddressCandidate.id))) == 1
        assert db.scalar(select(func.count(ApprovedAddress.id))) == 0


def _run_multisheet_job(client: TestClient) -> int:
    response = client.post(
        "/api/intake/upload/jobs",
        files={"file": ("multi_cex.xlsx", _multi_sheet_workbook(), "application/octet-stream")},
        data={"requested_source_type": "excel_upload"},
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["id"]
    run = client.post(f"/api/intake/jobs/{job_id}/run")
    assert run.status_code == 200, run.text
    assert run.json()["extracted_candidates"] == 2
    return job_id


def _multi_sheet_workbook() -> bytes:
    sheets = {
        "source_manifest": [
            [
                "sheet_name",
                "entity_hint",
                "source_url",
                "source_origin",
                "provenance_type",
                "evidence_shape",
                "snapshot_date",
                "operator_note",
            ],
            ["Bybit", "Bybit", "https://www.bybit.com/en/app/user/proof-of-reserve", "Bybit", "official_por_snapshot", "excel_wallet_list", "2026-04-22", "official Bybit sheet"],
            ["OKX", "OKX", "https://www.okx.com/proof-of-reserves", "OKX", "official_por_snapshot", "excel_wallet_list", "2026-04-22", "official OKX sheet"],
        ],
        "Bybit": [
            ["Entity", "Network", "Address", "Wallet Label / Role", "Evidence Type", "Confidence"],
            ["", "Ethereum", "0x1111111111111111111111111111111111111111", "Reserve Wallet", "excel_wallet_list", "90"],
        ],
        "OKX": [
            ["Entity", "Network", "Address", "Wallet Label / Role", "Evidence Type", "Confidence"],
            ["", "Ethereum", "0x2222222222222222222222222222222222222222", "Reserve Wallet", "excel_wallet_list", "90"],
        ],
    }
    return _xlsx_bytes(sheets)


def _xlsx_bytes(sheets: dict[str, list[list[str]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _job_with_candidate(db, *, entity_name: str, address: str) -> tuple[SourceJob, AddressCandidate]:
    preview = IntakePreview(
        id=str(uuid4()),
        source_artifact_json={},
        fingerprint_json={},
        profile_json={},
        preview_json={},
        warnings=[],
        fatal_errors=[],
    )
    db.add(preview)
    db.flush()
    job = SourceJob(
        preview_id=preview.id,
        input_method="upload",
        final_source_type="excel_upload",
        adapter_name="excel_csv_adapter",
        fingerprint_json={},
        source_artifact_json={},
        profile_json={"entity_name": entity_name},
        preview_json={},
        status="needs_review",
    )
    db.add(job)
    db.flush()
    document = SourceDocument(
        source_job_id=job.id,
        canonical_source_url=f"https://example.com/{entity_name.lower()}",
        file_path="source.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        document_title="source.xlsx",
        text_hash=str(uuid4()).replace("-", ""),
        metadata_json={},
    )
    db.add(document)
    db.flush()
    candidate = AddressCandidate(
        source_job_id=job.id,
        source_document_id=document.id,
        address=address,
        normalized_address=address.lower(),
        entity_name=entity_name,
        source_network="Ethereum",
        chain_guess="evm",
        chain_slug="ethereum",
        chain_id=1,
        address_family="evm",
        suggested_role="cex_por_wallet",
        confidence_initial=90,
        status="needs_review",
        source_type="excel_upload",
        source_input_type="xlsx_registry",
        source_sheet=entity_name,
        source_url=document.canonical_source_url,
        file_path=document.file_path,
        evidence_type="excel_wallet_list",
        warnings=[],
        raw_reference={"source_sheet": entity_name, "sheet_entity_hint": entity_name},
    )
    db.add(candidate)
    db.flush()
    db.add(
        AddressEvidence(
            candidate_id=candidate.id,
            source_document_id=document.id,
            evidence_type="excel_wallet_list",
            source_type="excel_upload",
            final_source_type="excel_upload",
            adapter_name="excel_csv_adapter",
            source_url=candidate.source_url,
            file_path=document.file_path,
            payload={"raw_reference": candidate.raw_reference},
            confidence_reason="structured_network_column",
        )
    )
    db.commit()
    db.refresh(job)
    db.refresh(candidate)
    return job, candidate


def _verify_candidate(db, candidate: AddressCandidate) -> None:
    record_source_verification(
        db,
        verification_scope="source_sheet",
        verification_status="verified",
        source_trust="official_verified",
        verified_by="pytest",
        source_job_id=candidate.source_job_id,
        source_document_id=candidate.source_document_id,
        source_sheet=candidate.source_sheet,
        entity_name=candidate.entity_name,
        source_url=candidate.source_url,
        evidence_shape=candidate.evidence_type,
    )
    db.commit()
